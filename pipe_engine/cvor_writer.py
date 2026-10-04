"""cvor_writer.py -- machine-readable Excel staging tables for CVOR XML.

The official delivery target for CVOR is XML/GGE, but Excel is still the
best engineering review surface.  These sheets are intentionally tabular:
each row maps cleanly to a future QuantityTakeoff-3_01 XML element.

The expanded ``ВР`` sheet remains the source of quantities and formulas.
CVOR sheets project that data into XML-shaped tables and add provenance,
catalog identifiers, document links, and validation checks.
"""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import PureWindowsPath
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .catalog_loader import CatalogBundle, CatalogRow


SHEET_VR = "ВР"
SHEET_CVOR_META = "ЦВОР_Meta"
SHEET_CVOR_FILES = "ЦВОР_Files"
SHEET_CVOR_SECTIONS = "ЦВОР_Sections"
SHEET_CVOR_WORKS = "ЦВОР_Works"
SHEET_CVOR_RESOURCES = "ЦВОР_Resources"
SHEET_CVOR_LINKS = "ЦВОР_Links"
SHEET_CVOR_SIGNATURES = "ЦВОР_Signatures"
SHEET_CVOR_CHECKS = "ЦВОР_Checks"

CVOR_SHEETS = [
    SHEET_CVOR_META,
    SHEET_CVOR_FILES,
    SHEET_CVOR_SECTIONS,
    SHEET_CVOR_WORKS,
    SHEET_CVOR_RESOURCES,
    SHEET_CVOR_LINKS,
    SHEET_CVOR_SIGNATURES,
    SHEET_CVOR_CHECKS,
]


@dataclass(frozen=True)
class _WorkRow:
    excel_row: int
    work_id: str
    section_id: str
    vr_row: int
    catalog_id: int | None
    item_type: str


def _catalog_id_from_note(note: str) -> int | None:
    marker = "catalog_id="
    if marker not in note:
        return None
    tail = note.split(marker, 1)[1]
    raw = tail.split(",", 1)[0].split(" ", 1)[0].strip()
    try:
        return int(raw)
    except ValueError:
        return None


def _catalog_by_id(catalog: CatalogBundle | None) -> dict[int, CatalogRow]:
    if catalog is None:
        return {}
    return {row.row_id: row for row in catalog.rows}


def _is_vr_work_row(vr, row: int) -> bool:
    item_num = vr.cell(row, 2).value
    return (
        isinstance(item_num, (int, float))
        and vr.cell(row, 4).value not in (None, "")
        and vr.cell(row, 5).value not in (None, "")
        and vr.cell(row, 6).value not in (None, "")
    )


def _is_vr_section_row(vr, row: int) -> bool:
    name = str(vr.cell(row, 4).value or "").strip()
    if not name:
        return False
    if _is_vr_work_row(vr, row):
        return False
    return name.endswith(":") or bool(vr.cell(row, 1).value)


def _guess_item_type(name: str, catalog_row: CatalogRow | None) -> str:
    low = name.lower()
    if any(word in low for word in ("разработка", "засыпка", "уплотнение", "вывоз", "устройство", "укладка", "протяжка")):
        return "работа"
    if catalog_row and catalog_row.tag:
        tag = catalog_row.tag.lower()
        if any(word in tag for word in ("грпш", "шкаф", "оборуд")):
            return "оборудование"
    return "материал" if catalog_row or name.startswith("Труба DN") else "прочее"


def _text_formula(value) -> str | None:
    """Store an Excel formula as text, not as an executable formula."""
    if value in (None, ""):
        return None
    text = str(value)
    if text.startswith("="):
        return "'" + text
    return text


def _write_headers(ws, headers: Iterable[str]) -> None:
    fill = PatternFill("solid", fgColor="D9E1F2")
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(1, col)
        cell.value = header
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 34


def _autosize(ws, widths: dict[int, int]) -> None:
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width


def build_cvor_sheets(
    wb: Workbook,
    *,
    vr_sheet_name: str = SHEET_VR,
    catalog: CatalogBundle | None = None,
    source_dwg_path: str | None = None,
    source_extracted_at_utc: str | None = None,
) -> None:
    """Create XML-shaped CVOR staging sheets from the generated ``ВР`` sheet."""
    if vr_sheet_name not in wb.sheetnames:
        return
    for sheet in CVOR_SHEETS:
        if sheet in wb.sheetnames:
            del wb[sheet]

    vr = wb[vr_sheet_name]
    catalog_rows = _catalog_by_id(catalog)

    file_id = "F001" if source_dwg_path else ""
    # DWG paths come from Windows CAD hosts: parse them as Windows paths on any OS.
    source_name = PureWindowsPath(source_dwg_path).name if source_dwg_path else ""

    _build_meta_sheet(wb, source_dwg_path=source_dwg_path, source_extracted_at_utc=source_extracted_at_utc)
    _build_files_sheet(wb, file_id=file_id, source_name=source_name, source_dwg_path=source_dwg_path)
    sections = _build_sections_sheet(wb, vr)
    works = _build_works_sheet(wb, vr, sections, catalog_rows)
    _build_resources_sheet(wb, works, catalog_rows)
    _build_links_sheet(wb, works, file_id=file_id, source_name=source_name)
    _build_signatures_sheet(wb)
    _build_checks_sheet(wb, works)


def _build_meta_sheet(
    wb: Workbook,
    *,
    source_dwg_path: str | None,
    source_extracted_at_utc: str | None,
) -> None:
    ws = wb.create_sheet(SHEET_CVOR_META)
    _write_headers(ws, ["Ключ", "Значение", "Комментарий"])
    rows = [
        ("schema_name", "QuantityTakeoff-3_01.xsd", "Целевая XSD-схема ВОР/ЦВОР"),
        ("target_extension", ".gge", "Финальный XML-файл для экспертизы"),
        ("construction_name", "", "Заполняется инженером/оркестратором проекта"),
        ("object_name", "", "Заполняется инженером/оркестратором проекта"),
        ("basis_document", "", "Основание разработки ВОР"),
        ("export_software", "pipe_engine", "Источник Excel-структуры"),
        ("source_dwg_path", source_dwg_path or "", "Исходный DWG из extract metadata"),
        ("source_extracted_at_utc", source_extracted_at_utc or "", "Время извлечения DWG"),
    ]
    for row, values in enumerate(rows, start=2):
        for col, value in enumerate(values, start=1):
            ws.cell(row, col).value = value
    _autosize(ws, {1: 28, 2: 72, 3: 56})


def _build_files_sheet(
    wb: Workbook,
    *,
    file_id: str,
    source_name: str,
    source_dwg_path: str | None,
) -> None:
    ws = wb.create_sheet(SHEET_CVOR_FILES)
    _write_headers(ws, ["FileID", "Тип", "Имя файла", "Путь", "Лист/страница", "Описание", "Хеш/версия"])
    if file_id:
        values = [file_id, "DWG", source_name, source_dwg_path, "Model", "Исходный чертеж трассы", ""]
        for col, value in enumerate(values, start=1):
            ws.cell(2, col).value = value
    else:
        ws.cell(2, 1).value = ""
        ws.cell(2, 6).value = "Добавить файл-основание перед XML export"
    _autosize(ws, {1: 12, 2: 10, 3: 34, 4: 84, 5: 16, 6: 42, 7: 18})


def _build_sections_sheet(wb: Workbook, vr) -> dict[int, str]:
    ws = wb.create_sheet(SHEET_CVOR_SECTIONS)
    _write_headers(ws, ["SectionID", "ParentID", "Номер", "Наименование", "VRRow"])
    row_to_section: dict[int, str] = {}
    current_section = "SEC-001"
    out_row = 2
    section_index = 1

    for vr_row in range(1, vr.max_row + 1):
        if _is_vr_section_row(vr, vr_row):
            section_id = f"SEC-{section_index:03d}"
            name = str(vr.cell(vr_row, 4).value or "").rstrip(":")
            ws.cell(out_row, 1).value = section_id
            ws.cell(out_row, 3).value = section_index
            ws.cell(out_row, 4).value = name
            ws.cell(out_row, 5).value = vr_row
            current_section = section_id
            section_index += 1
            out_row += 1
        row_to_section[vr_row] = current_section

    if out_row == 2:
        ws.cell(2, 1).value = current_section
        ws.cell(2, 3).value = 1
        ws.cell(2, 4).value = "Ведомость объемов работ"
        ws.cell(2, 5).value = ""
    _autosize(ws, {1: 14, 2: 14, 3: 10, 4: 56, 5: 10})
    return row_to_section


def _build_works_sheet(
    wb: Workbook,
    vr,
    row_to_section: dict[int, str],
    catalog_rows: dict[int, CatalogRow],
) -> list[_WorkRow]:
    ws = wb.create_sheet(SHEET_CVOR_WORKS)
    _write_headers(ws, [
        "WorkID",
        "SectionID",
        "Номер",
        "Type",
        "Name",
        "Unit",
        "Quantity",
        "QuantityFormula",
        "CatalogID",
        "Code",
        "TypeMark",
        "Manufacturer",
        "VRRow",
        "SourceDebug",
        "MetadataDebug",
    ])

    works: list[_WorkRow] = []
    out_row = 2
    for vr_row in range(1, vr.max_row + 1):
        if not _is_vr_work_row(vr, vr_row):
            continue
        note = str(vr.cell(vr_row, 9).value or "")
        catalog_id = _catalog_id_from_note(note)
        catalog_row = catalog_rows.get(catalog_id or -1)
        name = str(vr.cell(vr_row, 4).value or "")
        item_type = _guess_item_type(name, catalog_row)
        work_id = f"W{out_row - 1:05d}"

        values = [
            work_id,
            row_to_section.get(vr_row, "SEC-001"),
            f"='{SHEET_VR}'!C{vr_row}",
            item_type,
            f"='{SHEET_VR}'!D{vr_row}",
            f"='{SHEET_VR}'!E{vr_row}",
            f"='{SHEET_VR}'!F{vr_row}",
            _text_formula(vr.cell(vr_row, 8).value or vr.cell(vr_row, 6).value),
            catalog_id,
            catalog_row.code if catalog_row else None,
            catalog_row.type_mark if catalog_row else None,
            catalog_row.manufacturer if catalog_row else None,
            vr_row,
            f"='{SHEET_VR}'!G{vr_row}",
            f"='{SHEET_VR}'!I{vr_row}",
        ]
        for col, value in enumerate(values, start=1):
            ws.cell(out_row, col).value = value
        works.append(_WorkRow(out_row, work_id, row_to_section.get(vr_row, "SEC-001"), vr_row, catalog_id, item_type))
        out_row += 1

    _autosize(ws, {1: 12, 2: 12, 3: 12, 4: 14, 5: 60, 6: 12, 7: 14, 8: 42,
                   9: 12, 10: 18, 11: 34, 12: 28, 13: 10, 14: 28, 15: 44})
    return works


def _build_resources_sheet(
    wb: Workbook,
    works: list[_WorkRow],
    catalog_rows: dict[int, CatalogRow],
) -> None:
    ws = wb.create_sheet(SHEET_CVOR_RESOURCES)
    _write_headers(ws, [
        "ResourceID",
        "WorkID",
        "Type",
        "Name",
        "Unit",
        "Quantity",
        "CatalogID",
        "Code",
        "TypeMark",
        "Manufacturer",
    ])
    out_row = 2
    for work in works:
        if work.item_type not in ("материал", "оборудование") and work.catalog_id is None:
            continue
        catalog_row = catalog_rows.get(work.catalog_id or -1)
        values = [
            f"R{out_row - 1:05d}",
            work.work_id,
            work.item_type,
            f"='{SHEET_CVOR_WORKS}'!E{work.excel_row}",
            f"='{SHEET_CVOR_WORKS}'!F{work.excel_row}",
            f"='{SHEET_CVOR_WORKS}'!G{work.excel_row}",
            work.catalog_id,
            catalog_row.code if catalog_row else None,
            catalog_row.type_mark if catalog_row else None,
            catalog_row.manufacturer if catalog_row else None,
        ]
        for col, value in enumerate(values, start=1):
            ws.cell(out_row, col).value = value
        out_row += 1
    _autosize(ws, {1: 12, 2: 12, 3: 14, 4: 60, 5: 12, 6: 14, 7: 12, 8: 18, 9: 34, 10: 28})


def _build_links_sheet(
    wb: Workbook,
    works: list[_WorkRow],
    *,
    file_id: str,
    source_name: str,
) -> None:
    ws = wb.create_sheet(SHEET_CVOR_LINKS)
    _write_headers(ws, [
        "LinkID",
        "WorkID",
        "FileID",
        "PageNumber",
        "PageDescription",
        "SourceRef",
        "DWGHandle",
        "LayerOrBlock",
        "Comment",
    ])
    for out_row, work in enumerate(works, start=2):
        values = [
            f"L{out_row - 1:05d}",
            work.work_id,
            file_id,
            "",
            "Model" if file_id else "",
            f"='{SHEET_CVOR_WORKS}'!N{work.excel_row}",
            "",
            source_name,
            "Уточнить лист/страницу перед XML export" if file_id else "Нет FileID",
        ]
        for col, value in enumerate(values, start=1):
            ws.cell(out_row, col).value = value
    _autosize(ws, {1: 12, 2: 12, 3: 12, 4: 14, 5: 22, 6: 34, 7: 18, 8: 34, 9: 46})


def _build_signatures_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet(SHEET_CVOR_SIGNATURES)
    _write_headers(ws, ["SignatureID", "Role", "Organization", "Person", "Position", "Certificate/Note"])
    rows = [
        ("S001", "Разработал", "", "", "", ""),
        ("S002", "Проверил", "", "", "", ""),
    ]
    for row, values in enumerate(rows, start=2):
        for col, value in enumerate(values, start=1):
            ws.cell(row, col).value = value
    _autosize(ws, {1: 14, 2: 18, 3: 34, 4: 26, 5: 26, 6: 28})


def _build_checks_sheet(wb: Workbook, works: list[_WorkRow]) -> None:
    ws = wb.create_sheet(SHEET_CVOR_CHECKS)
    _write_headers(ws, ["Проверка", "Формула", "Статус", "Комментарий"])
    last = max((work.excel_row for work in works), default=1)
    checks = [
        ("Строк ВОР в ЦВОР", f"=COUNTA('{SHEET_CVOR_WORKS}'!A2:A{last})", "info", ""),
        ("Нет Type", f"=COUNTBLANK('{SHEET_CVOR_WORKS}'!D2:D{last})", "must_be_0", "Для XML Work.Type"),
        ("Нет QuantityFormula", f"=COUNTBLANK('{SHEET_CVOR_WORKS}'!H2:H{last})", "must_be_0", "Для доказуемости объема"),
        ("Нет FileID", f"=COUNTBLANK('{SHEET_CVOR_LINKS}'!C2:C{last})", "must_be_0", "Нужна ссылка на файл-основание"),
        ("Нет листа/страницы", f"=COUNTBLANK('{SHEET_CVOR_LINKS}'!D2:D{last})", "manual_review", "Для DWG Model допускается уточнить позже"),
        ("Материалы без кода", f'=COUNTIFS(\'{SHEET_CVOR_WORKS}\'!D2:D{last},"материал",\'{SHEET_CVOR_WORKS}\'!J2:J{last},"")', "review", "Коды есть не у всех строк каталога"),
        ("Нет подписантов", f"=COUNTBLANK('{SHEET_CVOR_SIGNATURES}'!D2:D3)", "manual_review", "Заполняется перед сдачей"),
    ]
    for row, values in enumerate(checks, start=2):
        for col, value in enumerate(values, start=1):
            ws.cell(row, col).value = value
    _autosize(ws, {1: 28, 2: 72, 3: 16, 4: 54})
