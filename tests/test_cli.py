"""
Tests for pipe_engine.cli — end-to-end smoke + argparse contract.

These tests invoke the CLI via the in-process `main()` entry point
(NOT via subprocess) so coverage/debugging still work cleanly.
"""
from __future__ import annotations

import io
import json
import tempfile
import unittest
from contextlib import redirect_stderr, redirect_stdout
from pathlib import Path

from openpyxl import load_workbook

from pipe_engine.cli import main as cli_main


FIXTURE = Path(__file__).resolve().parent / "fixtures" / "demo_extract.v1.json"


class TestCliBasic(unittest.TestCase):
    """End-to-end smoke: extract.json → CLI → valid xlsx with formulas."""

    def test_fill_writes_valid_xlsx(self):
        out_path = Path(tempfile.mkdtemp()) / "out.xlsx"
        rc = cli_main([
            "fill",
            "--extract", str(FIXTURE),
            "--output",  str(out_path),
            "--depth", "1.8",
            "--sand-pad", "0.1",
            "--snow",  "0.5",
            "--season", "winter",
            "--soil",   "sand",
            "--hand-trench-count", "5",
        ])
        self.assertEqual(rc, 0, "CLI should exit 0 on success")
        self.assertTrue(out_path.exists(), "xlsx should be created")
        self.assertGreater(out_path.stat().st_size, 2000, "xlsx should be >2KB")

        # Load back and verify it has both sheets
        wb = load_workbook(out_path, data_only=False)
        self.assertIn("ИсхДанные", wb.sheetnames)
        self.assertIn("Расчет земли", wb.sheetnames)

        # ИсхДанные!B1 = depth literal
        self.assertEqual(wb["ИсхДанные"]["B1"].value, 1.8)
        # ИсхДанные!B7 = winter flag = 1
        self.assertEqual(wb["ИсхДанные"]["B7"].value, 1)
        # ИсхДанные!B8 = hand_trench_count from CLI
        self.assertEqual(wb["ИсхДанные"]["B8"].value, 5)

        # Trace sheet has at least one data row (fixture has 2 main_g1 DNs)
        ws = wb["Расчет земли"]
        # Row 2 col 5 (E) = joint length formula
        v = ws.cell(2, 5).value
        self.assertIsInstance(v, str)
        self.assertTrue(v.startswith("="), f"E2 should be formula, got {v!r}")

    def test_fill_summer_disables_snow(self):
        """--season summer → ИсхДанные!B7 = 0 → C53 snow formula → 0."""
        out_path = Path(tempfile.mkdtemp()) / "summer.xlsx"
        rc = cli_main([
            "fill",
            "--extract", str(FIXTURE),
            "--output",  str(out_path),
            "--season", "summer",
        ])
        self.assertEqual(rc, 0)
        wb = load_workbook(out_path, data_only=False)
        self.assertEqual(wb["ИсхДанные"]["B7"].value, 0,
                         "season=summer should set ИсхДанные!B7 = 0")

    def test_fill_sand_cushion_flag(self):
        """--sand-cushion → ИсхДанные!B6 = 1."""
        out_path = Path(tempfile.mkdtemp()) / "cushion.xlsx"
        rc = cli_main([
            "fill",
            "--extract", str(FIXTURE),
            "--output",  str(out_path),
            "--sand-cushion",
        ])
        self.assertEqual(rc, 0)
        wb = load_workbook(out_path, data_only=False)
        self.assertEqual(wb["ИсхДанные"]["B6"].value, 1,
                         "--sand-cushion should set ИсхДанные!B6 = 1")

    def test_fill_missing_extract_returns_error(self):
        out_path = Path(tempfile.mkdtemp()) / "x.xlsx"
        with redirect_stderr(io.StringIO()):
            rc = cli_main([
                "fill",
                "--extract", "/nonexistent/path.json",
                "--output",  str(out_path),
            ])
        self.assertEqual(rc, 1, "missing extract should return rc=1")
        self.assertFalse(out_path.exists(), "no xlsx should be created on error")

    def test_fill_bad_json_returns_error(self):
        bad = Path(tempfile.mkdtemp()) / "bad.json"
        bad.write_text("{not valid json at all", encoding="utf-8")
        out_path = Path(tempfile.mkdtemp()) / "x.xlsx"
        with redirect_stderr(io.StringIO()):
            rc = cli_main([
                "fill",
                "--extract", str(bad),
                "--output",  str(out_path),
            ])
        self.assertEqual(rc, 1, "bad JSON should return rc=1")


class TestCliHelp(unittest.TestCase):
    def test_top_help(self):
        buf = io.StringIO()
        with redirect_stdout(buf):
            with self.assertRaises(SystemExit) as ctx:
                cli_main(["--help"])
        # argparse --help exits with code 0
        self.assertEqual(ctx.exception.code, 0)
        self.assertIn("pipe-engine", buf.getvalue())

    def test_fill_help(self):
        buf = io.StringIO()
        with redirect_stdout(buf):
            with self.assertRaises(SystemExit) as ctx:
                cli_main(["fill", "--help"])
        self.assertEqual(ctx.exception.code, 0)
        out = buf.getvalue()
        # Sanity: all major flags appear in help
        for flag in ["--extract", "--output", "--depth", "--soil", "--season",
                     "--sand-cushion", "--hand-trench-count"]:
            self.assertIn(flag, out, f"{flag!r} should appear in fill --help")


if __name__ == "__main__":
    unittest.main()
